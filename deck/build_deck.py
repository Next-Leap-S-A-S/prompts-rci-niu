#!/usr/bin/env python3
"""Genera la presentacion del modulo 3 con el estilo visual de la plantilla de NIU.

Usa como base el PPTX de referencia (formas, fondos y tipografias incrustadas) y lo
reemplaza por las laminas nuevas. Uso:
  python3 deck/build_deck.py "<plantilla.pptx>" "<salida.pptx>"
"""
import copy
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
TITLE, BOLD, BODY = "Codec Pro Ultra-Bold", "Codec Pro Bold", "Codec Pro"
WHITE, GOLD, MAROON, NAVY, INK, PAPER = "FFFFFF", "F5A81C", "801711", "15193E", "111114", "E4E5EC"
LOGO_BLANCO = sys.argv[3] if len(sys.argv) > 3 else None

prs = Presentation(SRC)
orig = list(prs.slides)
blank = orig[0].slide_layout


def src(n):
    return orig[n - 1]


DECOR = {
    "cover": (src(1), ["Freeform %d" % i for i in range(2, 10)]),
    "dark": (src(8), ["Freeform %d" % i for i in range(2, 9)]),
    "light": (src(3), ["Freeform %d" % i for i in range(2, 10)]),
}


def clone_decor(kind, dst):
    s, names = DECOR[kind]
    bg = s._element.cSld.find(qn("p:bg"))
    if bg is not None:
        dst._element.cSld.insert(0, copy.deepcopy(bg))
    for sh in s.shapes:
        if sh.name not in names:
            continue
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for k, v in list(node.attrib.items()):
                if k.endswith("}embed") or k.endswith("}link"):
                    rel = s.part.rels[v]
                    node.set(k, dst.part.relate_to(rel.target_part, rel.reltype))
        dst.shapes._spTree.append(el)


def rgb(h):
    return RGBColor.from_string(h)


def alpha(shape_fill, pct):
    clr = shape_fill._xPr.find(qn("a:solidFill"))[0]
    a = clr.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))})
    clr.append(a)


def tb(slide, x, y, w, h, paras, size=24, color=WHITE, font=BODY, bold=False, align="l", anchor="t", spacing=1.12):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        para.line_spacing = spacing
        runs = p if isinstance(p, list) else [(p, font, color, bold)]
        for r in runs:
            text, f, c, b = r if isinstance(r, tuple) else (r, font, color, bold)
            run = para.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.name = f
            run.font.bold = False
            run.font.color.rgb = rgb(c)
    return box


def card(slide, x, y, w, h, fill=NAVY, a=62, line=None, radius=0.12):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    alpha(s.fill, a)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def pill(slide, x, y, w, h, text, size=26, fill=MAROON, color=WHITE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.5
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = BOLD
    r.font.color.rgb = rgb(color)
    return s


def new(kind):
    s = prs.slides.add_slide(blank)
    clone_decor(kind, s)
    return s


def headline(s, text, size=58, y=2.35, w=14.0, color=WHITE, x=1.4):
    return tb(s, x, y, w, 2.2, text, size=int(size * 0.92), font=TITLE, color=color, spacing=1.02)


def cards(s, items, y=4.6, h=4.6, cols=None, x0=1.4, W=17.2, gap=0.4, tsize=28, bsize=21, tag=None):
    n = len(items)
    y += 0.55
    h -= 0.45
    tsize = int(tsize * 1.08)
    bsize = int(bsize * 1.2)
    cols = cols or n
    w = (W - gap * (cols - 1)) / cols
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = x0 + c * (w + gap)
        yy = y + r * (h + gap)
        card(s, x, yy, w, h)
        title, body = it[0], it[1]
        tb(s, x + 0.35, yy + 0.3, w - 0.7, 1.2, title, size=tsize, font=BOLD, color=GOLD, spacing=1.0)
        tb(s, x + 0.35, yy + 0.3 + (1.35 if len(title) > 22 and tsize >= 26 else 0.95), w - 0.7, h - 1.6, body, size=bsize, color=WHITE, spacing=1.15)


# ---------------- LAMINAS ----------------
# 1 portada
s = new("cover")
if LOGO_BLANCO:
    s.shapes.add_picture(LOGO_BLANCO, Inches(7.9), Inches(1.7), height=Inches(1.55))
tb(s, 1.0, 3.9, 18, 3.4, ["PROMPT ENGINEERING", "en el derecho"], size=92, font=TITLE, align="c", spacing=1.0)
pill(s, 3.2, 7.75, 13.6, 1.1, "MÓDULO 3 · HERRAMIENTAS AVANZADAS Y GPTs PERSONALIZADOS", size=24)
tb(s, 3.2, 9.35, 13.6, 0.6, "Diplomado en IA y Protección de Datos · Miguel Martínez · Next Leap", size=20, align="c", color="E4E5EC")

# 2 objetivos (clara)
s = new("light")
tb(s, 1.2, 1.5, 11.0, 1.4, "OBJETIVOS", size=64, font=TITLE, color="000000", align="c")
tb(s, 1.6, 3.2, 10.4, 6.0,
   ["1. Escribir prompts jurídicos con el método Rol, Contexto, Instrucción.",
    "2. Elegir la herramienta adecuada para investigar, argumentar y analizar.",
    "3. Diseñar un asistente personalizado para una tarea jurídica repetida.",
    "4. Reconocer y mitigar alucinaciones, complacencia y sesgos.",
    "5. Aplicar la supervisión humana y la protección de datos personales."],
   size=29, color=INK, spacing=1.12)
tb(s, 13.6, 3.6, 6.2, 4.2, ["Hoy construimos un prompt, lo probamos en varios modelos y lo ponemos en duda."], size=40, font=BOLD, color=WHITE, spacing=1.2)

# 3 evaluacion
s = new("dark")
headline(s, "Cómo se evalúa el módulo", size=56)
stats = [("7/10", "Nota mínima para aprobar el módulo."), ("30%", "Tres tareas semanales, 10% cada una."), ("70%", "Tarea final: laboratorio de prompts.")]
w = (17.2 - 0.8) / 3
for i, (n, t) in enumerate(stats):
    x = 1.4 + i * (w + 0.4)
    card(s, x, 5.0, w, 4.3)
    tb(s, x + 0.4, 5.3, w - 0.8, 1.9, n, size=104, font=TITLE, color=GOLD)
    tb(s, x + 0.4, 7.6, w - 0.8, 1.5, t, size=30, color=WHITE)

# 4 ruta
s = new("dark")
headline(s, "Cuatro unidades, un hilo conductor", size=56)
cards(s, [("U1 · Prompting y método RCI", "Cómo pedir trabajo jurídico a un modelo sin que invente."),
          ("U2 · Modelos para investigar", "NotebookLM, proyectos y búsqueda profunda con fuentes."),
          ("U3 · Diseño de asistentes", "GPTs y gems con instrucciones y documentos propios."),
          ("U4 · Riesgos y mitigación", "Alucinación, complacencia, sesgos y datos personales.")],
      y=4.4, h=4.6, cols=4, tsize=24, bsize=21)


def section(num, title, sub):
    s = new("cover")
    tb(s, 1.5, 2.6, 17, 1.5, "UNIDAD %d" % num, size=40, font=BOLD, color=GOLD, align="c")
    tb(s, 1.5, 4.0, 17, 3.2, title, size=76, font=TITLE, align="c", spacing=1.02)
    pill(s, 4.0, 7.9, 12.0, 1.05, sub, size=24)


# U1
section(1, "Introducción al prompting y metodología RCI", "Cómo pedirle trabajo jurídico a un modelo")

s = new("dark")
headline(s, "Un prompt vago produce un escrito que el juez inadmite.", size=60)
card(s, 1.4, 5.3, 17.2, 2.2, fill=MAROON, a=70)
tb(s, 1.9, 5.65, 16.2, 1.6, "«Hazme una demanda contra un proveedor que no cumplió.»", size=36, font=BOLD, color=WHITE, anchor="m")
tb(s, 1.4, 8.1, 17.2, 1.8, "No dice qué contrato, qué incumplimiento ni qué se pide. El modelo inventa los hechos y las pretensiones que faltan.", size=30, color="E4E5EC")

s = new("dark")
headline(s, "Tres partes. Un solo método.", size=62)
cards(s, [("ROL", "Quién debería hacer la tarea. Formación, años de práctica, sesgo útil y tono."),
          ("CONTEXTO", "Los hechos, el estado del trámite, los documentos y lo que ya se intentó."),
          ("INSTRUCCIÓN", "Pasos numerados, formato de salida, límites y criterios de verificación.")],
      y=4.4, h=4.8, cols=3, tsize=34, bsize=25)

s = new("dark")
headline(s, "El mismo pedido, con método", size=56)
card(s, 1.4, 4.8, 6.7, 4.9, fill=MAROON, a=70)
tb(s, 1.8, 5.1, 5.9, 0.7, "ANTES", size=26, font=BOLD, color=GOLD)
tb(s, 1.8, 5.9, 5.9, 3.6, "«Revisa esta política de privacidad y dime si cumple la ley.»", size=34, font=BOLD)
card(s, 8.4, 4.8, 10.2, 4.9)
tb(s, 8.8, 5.1, 9.4, 0.7, "DESPUÉS", size=26, font=BOLD, color=GOLD)
tb(s, 8.8, 5.9, 9.4, 3.7,
   [[("Rol. ", BOLD, GOLD, False), ("Abogado de protección de datos con 10 años de experiencia ante la LOPDP.", BODY, WHITE, False)],
    [("Contexto. ", BOLD, GOLD, False), ("Una aplicación en Ecuador, qué datos recoge, para qué y con quién los comparte.", BODY, WHITE, False)],
    [("Instrucción. ", BOLD, GOLD, False), ("Tabla de finalidades y bases, máximo 700 palabras, cada norma con [verificar], sin concluir cumplimiento.", BODY, WHITE, False)]],
   size=26, spacing=1.2)

s = new("dark")
headline(s, "Cuatro técnicas para afinar el prompt", size=56)
cards(s, [("Zero-shot", "Se pide sin ejemplos. Sirve para tareas simples y bien definidas."),
          ("Few-shot", "Se incluyen dos o tres ejemplos del formato esperado. El modelo imita su estructura."),
          ("Prompting negativo", "Se dice qué no hacer: «no inventes sentencias», «no concluyas»."),
          ("Verificación cruzada", "Se pide listar lo que debe comprobarse y con qué fuente.")],
      y=4.4, h=4.8, cols=4, tsize=26, bsize=21)

s = new("dark")
headline(s, "Del ensayo y error al diseño.", size=64)
tb(s, 1.4, 4.9, 15.5, 4.6,
   ["Un prompt que funcionó es una plantilla. Se guarda y solo se cambian los corchetes.",
    "Casi todo el resultado depende de tres decisiones: el rol, el contexto concreto y el formato de salida. Empiece por ahí."],
   size=34, spacing=1.3)

s = new("dark")
headline(s, "Actividad: del prompt vago al prompt RCI", size=52)
cards(s, [("1. Escriba", "Su prompt vago para una tarea real de su práctica."),
          ("2. Reescriba", "En Rol, Contexto, Instrucción. Observe el medidor del laboratorio."),
          ("3. Pruebe", "Ejecute ambas versiones en dos modelos y compare."),
          ("4. Entregue", "Es la Tarea 1 y vale el 10% de la nota.")],
      y=4.4, h=4.6, cols=4, tsize=28, bsize=22)

# U2
section(2, "Modelos de IA generativa para investigación y análisis", "Elegir la herramienta según la tarea")

s = new("dark")
headline(s, "Tres herramientas, tres usos", size=60)
cards(s, [("NotebookLM", "Trabaja solo con las fuentes que usted carga y señala de dónde sale cada respuesta."),
          ("Proyectos", "En Claude y ChatGPT guardan instrucciones y archivos del asunto para todas las conversaciones."),
          ("Búsqueda profunda", "Investiga en la web de forma iterativa y entrega un informe con enlaces. Cada enlace se abre y se verifica.")],
      y=4.4, h=4.8, cols=3, tsize=32, bsize=24)

s = new("dark")
headline(s, "Si no hay fuente, el sistema debe decir que no sabe.", size=54)
cards(s, [("1. La base", "Se construyen los documentos a los que se conectará la IA."),
          ("2. La búsqueda", "El modelo selecciona las fuentes que sirven. Si no hay, lo dice."),
          ("3. La respuesta", "Redacta con base en las fuentes y permite rastrear cada dato.")],
      y=4.6, h=3.6, cols=3, tsize=28, bsize=23)
tb(s, 1.4, 8.9, 17.2, 1.6, "Caja negra: no se ve en qué se apoyó. Caja de cristal: se rastrea hasta el documento. En derecho se prefiere la segunda.", size=28, color="E4E5EC")

s = new("dark")
headline(s, "Lo que la IA lee completo", size=58)
cards(s, [("Texto sin estructura", ".txt, .md, .pdf, .docx. El texto plano pesa poco y deja leer solo el contenido."),
          ("Datos estructurados", ".csv y .xlsx. Útiles para cuantías, términos y listados de procesos."),
          ("Escaneos y fotos", "Requieren reconocimiento óptico de caracteres antes de cargarse. Sin eso, el modelo no lee el texto.")],
      y=4.4, h=4.0, cols=3, tsize=28, bsize=22)
tb(s, 1.4, 9.1, 17.2, 1.6, "La ventana de contexto incluye las instrucciones, los archivos, parte de la conversación y el historial del chat.", size=28, color="E4E5EC")

# U3
section(3, "Diseño de asistentes con IA generativa", "De un prompt repetido a un asistente propio")

s = new("dark")
headline(s, "Chat suelto, asistente, agente", size=58)
cards(s, [("Chat suelto", "Se escribe el prompt cada vez. No recuerda su criterio."),
          ("Asistente personalizado", "Guarda instrucciones y documentos. Repite una tarea con el mismo criterio."),
          ("Agente", "Además actúa: busca, ejecuta y encadena pasos. Exige límites más estrictos.")],
      y=4.4, h=3.9, cols=3, tsize=28, bsize=22)
tb(s, 1.4, 9.0, 17.2, 1.7, "En ChatGPT se llama GPT, en Claude proyecto, en Gemini gem y en Microsoft 365 Copilot agente. La receta se pega igual en todos.", size=28, color="E4E5EC")

s = new("dark")
headline(s, "Siete ingredientes", size=62)
ing = [("1 · Nombre", "Qué hace y para quién."), ("2 · Instrucciones", "El prompt RCI, escrito una vez."), ("3 · Modelo", "El más capaz para razonar."),
       ("4 · Documentos de apoyo", "Lo que no cambia entre casos."), ("5 · Archivos del caso", "Lo que sí cambia. Anonimizado."),
       ("6 · Herramientas", "Cada una amplía lo que puede salir mal."), ("7 · Límites", "Qué no concluye y cuándo escala.")]
w = (17.2 - 0.3 * 3) / 4
for i, (t, b) in enumerate(ing):
    r, c = divmod(i, 4)
    x, y = 1.4 + c * (w + 0.3), 4.8 + r * 2.35
    card(s, x, y, w, 2.1)
    tb(s, x + 0.3, y + 0.2, w - 0.6, 0.9, t, size=25, font=BOLD, color=GOLD)
    tb(s, x + 0.3, y + 1.0, w - 0.6, 1.1, b, size=21)

s = new("dark")
headline(s, "Un asistente listo: gestor de solicitudes de titulares", size=48)
card(s, 1.4, 4.9, 8.3, 4.7)
tb(s, 1.8, 5.2, 7.5, 0.7, "INSTRUCCIÓN, EN BREVE", size=24, font=BOLD, color=GOLD)
tb(s, 1.8, 6.0, 7.5, 3.4, "Clasifica la solicitud, indica el plazo y pide verificarlo. Analiza si procede total o parcialmente. Redacta la respuesta al titular y una nota para el expediente.", size=26, spacing=1.2)
card(s, 10.0, 4.9, 8.6, 4.7, fill=MAROON, a=70)
tb(s, 10.4, 5.2, 7.8, 0.7, "SU LÍMITE", size=24, font=BOLD, color=GOLD)
tb(s, 10.4, 6.0, 7.8, 3.4, ["No decide si se niega la solicitud. Esa decisión es del abogado responsable.", "No entrega datos de terceros.", "Termina siempre con: «Respuesta sujeta a revisión del abogado responsable»."], size=26, spacing=1.2)

s = new("dark")
headline(s, "Actividad: construya un asistente y póngalo a prueba", size=48)
cards(s, [("1. Configure", "Pegue una receta en su plataforma y cargue un documento de apoyo ficticio."),
          ("2. Ejecute", "Corra los tres mensajes de prueba de la receta."),
          ("3. Rompa sus reglas", "Pegue una cédula, pida una conclusión que no le toca, pida una sentencia que respalde su postura."),
          ("4. Entregue", "Es la Tarea 3 y vale el 10% de la nota.")],
      y=4.4, h=5.0, cols=4, tsize=26, bsize=21)

# U4
section(4, "Riesgos en el ámbito jurídico y técnicas de mitigación", "Lo que puede salir mal y cómo evitarlo")

s = new("dark")
headline(s, "Cuatro trampas conocidas", size=62)
cards(s, [("Alucinación", "Inventa normas, artículos y sentencias con apariencia verosímil. Se mitiga abriendo cada cita en la fuente oficial."),
          ("Complacencia", "Le da la razón a quien pregunta. Se mitiga pidiendo la posición contraria y los puntos débiles."),
          ("Sesgos", "Reproduce los de sus datos de entrenamiento. Se mitiga con criterios objetivos de evaluación."),
          ("Datos personales", "Lo que se pega puede salir del control del despacho. Se mitiga anonimizando antes de pegar.")],
      y=4.4, h=5.2, cols=4, tsize=26, bsize=20)

s = new("dark")
headline(s, "Mata v. Avianca, 2023", size=64)
tb(s, 1.4, 4.8, 11.5, 5.0,
   ["Dos abogados presentaron ante un tribunal de Nueva York escritos con sentencias que no existían. Las había generado ChatGPT.",
    "El juez los sancionó con 5.000 dólares. Cuando se les advirtió el error, insistieron en las citas.",
    "La lección es una regla de oficio: nada que salga de la IA se cita sin abrir la fuente."],
   size=28, spacing=1.25)
card(s, 13.6, 4.8, 5.0, 4.6, fill=MAROON, a=75)
tb(s, 13.9, 5.2, 4.4, 1.6, "5.000", size=70, font=TITLE, color=GOLD)
tb(s, 13.9, 6.9, 4.4, 2.4, "dólares de sanción por citas inexistentes", size=28)

s = new("dark")
headline(s, "Antes de pegar el caso", size=62)
tb(s, 1.4, 4.9, 17, 5.2,
   ["Reemplace por marcadores cédulas, RUC, nombres, números de causa, placas, correos y teléfonos.",
    "Use solo las herramientas que su organización haya autorizado.",
    "Si el tratamiento es nuevo, haga una evaluación de impacto con la LOPDP.",
    "Cada norma que cite la IA se marca [verificar] y se confirma en la fuente."],
   size=32, spacing=1.3)

s = new("dark")
tb(s, 1.5, 3.2, 17, 3.0, "Supervisión humana significativa.", size=84, font=TITLE, align="c", spacing=1.0)
tb(s, 3.0, 7.4, 14, 1.8, "La IA redacta. El abogado decide, verifica y firma.", size=36, font=BOLD, align="c", color=GOLD)

# tarea final
s = new("dark")
headline(s, "Tarea final: laboratorio de prompts", size=56)
steps = [("1", "Un prompt en Rol, Contexto, Instrucción con tres iteraciones"), ("2", "Tres modelos, tres casos: normal, adversarial y con datos personales"),
         ("3", "Análisis comparativo entre modelos"), ("4", "Evaluación de impacto y protocolo con la LOPDP"), ("5", "Sustentación oral de diez minutos")]
w = (17.2 - 0.3 * 4) / 5
for i, (n, t) in enumerate(steps):
    x = 1.4 + i * (w + 0.3)
    card(s, x, 4.9, w, 3.9)
    tb(s, x + 0.3, 5.1, w - 0.6, 1.3, n, size=60, font=TITLE, color=GOLD)
    tb(s, x + 0.3, 6.4, w - 0.6, 2.3, t, size=23)
tb(s, 1.4, 9.1, 17.2, 1.5, "Vale el 70%. Ficha 10%, prompt 25%, pruebas 25%, análisis comparativo 20%, impacto y protocolo 10%, sustentación 10%.", size=24, color="E4E5EC")

s = new("dark")
headline(s, "Qué sigue en el aula virtual", size=58)
cards(s, [("Material", "Página del módulo y laboratorio de prompts, para descargar y abrir en el navegador."),
          ("Tareas", "Tres tareas semanales y la tarea final, con objetivos, formato y fecha límite."),
          ("Test", "Diez preguntas de autoevaluación con retroalimentación inmediata.")],
      y=4.4, h=4.2, cols=3, tsize=32, bsize=24)

# cierre
s = new("cover")
if LOGO_BLANCO:
    s.shapes.add_picture(LOGO_BLANCO, Inches(7.9), Inches(1.9), height=Inches(1.55))
tb(s, 1.5, 4.1, 17, 3.2, ["El prompt es una herramienta.", "La responsabilidad es suya."], size=76, font=TITLE, align="c", spacing=1.02)
pill(s, 5.2, 8.0, 9.6, 1.0, "Miguel Martínez · Next Leap", size=26)

# ---- quitar laminas originales ----
lst = prs.slides._sldIdLst
for sid in list(lst)[: len(orig)]:
    prs.part.drop_rel(sid.rId)
    lst.remove(sid)

prs.save(OUT)
print("ok", OUT, len(prs.slides), "laminas")
